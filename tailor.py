"""Per-offer CV tailoring: Claude adapts the Word CV, code enforces that missing technologies never become experience."""

from __future__ import annotations

import copy
import re
import unicodedata
from pathlib import Path

import llm
import radar
from radar import normalize

OUT = radar.ROOT / "data" / "cv_tailored"
DEFAULTS = {
    "enabled": True,
    "tailor_summary": True,
    "tailor_experience": True,
    "add_missing_skills": True,
    "max_added_skills": 6,
    "missing_in_cover_letter": True,
    "rules": ("Adaptá el CV al aviso usando sus mismas palabras cuando describan algo que realmente hice. "
              "Reordená viñetas y habilidades para que lo más relevante quede primero. "
              "Nunca inventes logros, cifras, empresas, fechas ni responsabilidades."),
    "cover_letter_rules": ("Carta breve (4 a 7 oraciones), concreta y sin frases infladas. Mencioná 2 o 3 coincidencias "
                           "reales entre el aviso y mi experiencia. Las tecnologías que pide el aviso y no están en mi "
                           "experiencia nombralas como tecnologías que manejo o con las que estoy trabajando, nunca como "
                           "experiencia laboral ni con años de uso. No te disculpes ni señales carencias: no menciones requisitos "
                           "que no cumplo (seniority, industria, años), salvo nombrar las tecnologías faltantes así."),
}


def settings(cfg: dict) -> dict:
    return {**DEFAULTS, **cfg.get("tailoring", {})}


# ---------- reading the Word CV ----------

def _is_list(p) -> bool:
    ppr = p._p.pPr
    return ppr is not None and ppr.numPr is not None


def _is_heading(p) -> bool:
    text = p.text.strip()
    return bool(text) and not _is_list(p) and text == text.upper() and bool(p.runs) and bool(p.runs[0].bold)


def parse(document) -> dict:
    """Locate summary, skills ('Label: value' bullets) and experience roles in the CV."""
    sections, current = {}, None
    for p in document.paragraphs:
        if _is_heading(p) and p is not document.paragraphs[0]:
            key = normalize(p.text)
            current = ("summary" if any(k in key for k in ("resumen", "summary", "perfil", "profile")) else
                       "skills" if any(k in key for k in ("habilidades", "skills", "tecnolog")) else
                       "experience" if any(k in key for k in ("experiencia", "experience")) else "other")
            sections.setdefault(current, [])
            continue
        if current and p.text.strip():
            sections[current].append(p)
    roles = []
    for p in sections.get("experience", []):
        if _is_list(p):
            if roles:
                roles[-1]["bullets"].append(p)
        else:
            roles.append({"header": p, "bullets": []})
    skills = [p for p in sections.get("skills", []) if _is_list(p) and ":" in p.text]
    return {"summary": sections.get("summary", [None])[0], "skills": skills, "roles": roles}


def split_items(value: str) -> list[str]:
    """Split 'Python, SQL (T-SQL, MySQL), Git' on commas outside parentheses."""
    items, depth, buf = [], 0, ""
    for ch in value:
        depth += ch == "("
        depth -= ch == ")"
        if ch == "," and depth == 0:
            items.append(buf.strip())
            buf = ""
        else:
            buf += ch
    if buf.strip():
        items.append(buf.strip())
    return [i for i in items if i]


def base_facts(parsed: dict) -> dict:
    skills = [{"label": p.text.split(":", 1)[0].strip(), "value": p.text.split(":", 1)[1].strip()} for p in parsed["skills"]]
    return {"summary": parsed["summary"].text if parsed["summary"] else "",
            "skills": skills,
            "roles": [{"header": r["header"].text.strip(), "bullets": [b.text.strip() for b in r["bullets"]]}
                      for r in parsed["roles"]]}


# ---------- asking Claude ----------

SCHEMA = {
    "type": "object",
    "properties": {
        "summary": {"type": "string"},
        "skills": {"type": "array", "items": {"type": "object", "properties": {
            "label": {"type": "string"}, "value": {"type": "string"}}, "required": ["label", "value"]}},
        "roles": {"type": "array", "items": {"type": "object", "properties": {
            "bullets": {"type": "array", "items": {"type": "string"}}}, "required": ["bullets"]}},
        "missing": {"type": "array", "items": {"type": "string"},
                    "description": "Tecnologías o requisitos del aviso que no figuran en la experiencia del CV"},
        "cover_letter": {"type": "string"},
    },
    "required": ["summary", "skills", "roles", "missing", "cover_letter"],
}


def ask_claude(job: dict, facts: dict, cfg: dict, language: str, as_email: bool) -> dict:
    s = settings(cfg)
    import agent  # local import: agent imports this module
    roles = "\n".join(f"[{i}] {r['header']}\n" + "\n".join(f"  - {b}" for b in r["bullets"])
                      for i, r in enumerate(facts["roles"]))
    skills = "\n".join(f"- {x['label']}: {x['value']}" for x in facts["skills"])
    prompt = (
        "Adaptá mi CV a esta oferta y escribí la carta de presentación.\n\n"
        "REGLAS DURAS (se verifican automáticamente y lo que las viole se descarta):\n"
        "1. Resumen y viñetas de experiencia: solo hechos que ya están en el CV. Podés reescribirlos, reordenarlos y usar "
        "el vocabulario del aviso, pero NO agregues tecnologías, herramientas ni tareas que no figuren en esa misma viñeta "
        "o en el CV. Mantené la cantidad de roles y no agregues viñetas.\n"
        + ("2. Habilidades técnicas: podés agregar las tecnologías que pide el aviso y no tengo (máximo "
           f"{s['max_added_skills']}), ubicándolas en la categoría que corresponda o en una línea nueva "
           "«Otras tecnologías» / «Other technologies». Conservá todas las categorías y valores existentes, reordenados "
           "por relevancia.\n" if s["add_missing_skills"] else
           "2. Habilidades técnicas: solo reordená; no agregues tecnologías nuevas.\n")
        + ("" if s["tailor_summary"] else "3. Devolvé el resumen exactamente igual.\n")
        + ("" if s["tailor_experience"] else "4. Devolvé las viñetas exactamente iguales.\n")
        + f"\nREGLAS DE ESTILO DEL CV:\n{s['rules']}\n\n"
        f"CARTA ({'cuerpo del correo con saludo y firma, el CV va adjunto' if as_email else 'texto para pegar en el formulario, sin firma de correo'}):\n"
        f"{s['cover_letter_rules']}"
        + ("" if s["missing_in_cover_letter"] else "\nNo menciones en la carta las tecnologías que no tengo.")
        + f"\n\nIdioma de todo el contenido: {'inglés' if language == 'en' else 'español'}.\n\n"
        f"=== MI CV ===\nRESUMEN:\n{facts['summary']}\n\nHABILIDADES:\n{skills}\n\nEXPERIENCIA:\n{roles}\n\n"
        f"=== OFERTA ===\nCargo: {job['title']}\nEmpresa: {job['company']}\nUbicación: {job['location']}\n\n"
        f"{(job.get('description') or '')[:8000]}")
    return llm.ask(agent.system_prompt(cfg), prompt, SCHEMA, model=cfg["autopilot"].get("model", "sonnet"))


# ---------- enforcing the rules ----------

def guard(facts: dict, proposal: dict, cfg: dict) -> tuple[dict, list[str], list[str]]:
    """Revert any experience/summary text that introduces technologies absent from the CV.
    Returns (clean proposal, added skills, reverted notes)."""
    s = settings(cfg)
    original_text = facts["summary"] + " " + " ".join(f"{x['label']} {x['value']}" for x in facts["skills"]) + " " + \
        " ".join(b for r in facts["roles"] for b in r["bullets"])
    known = set(radar.technologies(original_text))
    reverted = []

    summary = proposal["summary"].strip() if s["tailor_summary"] else facts["summary"]
    extra = set(radar.technologies(summary)) - known
    if extra or not summary:
        if extra:
            reverted.append(f"Resumen: mencionaba {', '.join(sorted(extra))}; se dejó el original")
        summary = facts["summary"]

    roles = []
    for i, base in enumerate(facts["roles"]):
        new = proposal["roles"][i]["bullets"] if s["tailor_experience"] and i < len(proposal["roles"]) else base["bullets"]
        new = [b.strip() for b in new if b.strip()][:len(base["bullets"])]
        if len(new) < len(base["bullets"]):
            new = base["bullets"]
        extra = set(radar.technologies(" ".join(new))) - known
        if extra:
            # Bullets may be reordered, so a violation reverts the whole role to its original wording.
            reverted.append(f"{base['header'][:40]}: mencionaba {', '.join(sorted(extra))}; se dejaron las viñetas originales")
            new = base["bullets"]
        roles.append({"bullets": new})

    # Skills: keep every original item, allow up to N new ones.
    original_items = {normalize(i) for x in facts["skills"] for i in split_items(x["value"])}
    import agent  # local import: agent imports this module
    rule = agent.cloud_filter(cfg)
    blocked_cloud = rule["blocked_terms"] if rule["enabled"] else []
    added, skills = [], []
    for x in proposal["skills"]:
        keep = []
        for item in split_items(x["value"]):
            if normalize(item) in original_items:
                keep.append(item)
            elif agent._terms_in(item, blocked_cloud):
                continue  # Microsoft-only search: never add AWS/GCP skills to the CV
            elif s["add_missing_skills"] and len(added) < s["max_added_skills"]:
                keep.append(item)
                added.append(item)
        if keep:
            skills.append({"label": x["label"].strip(), "value": ", ".join(keep)})
    present = {normalize(i) for x in skills for i in split_items(x["value"])}
    for x in facts["skills"]:  # anything Claude dropped goes back in its original line
        lost = [i for i in split_items(x["value"]) if normalize(i) not in present]
        if lost:
            line = next((y for y in skills if normalize(y["label"]) == normalize(x["label"])), None)
            if line:
                line["value"] += ", " + ", ".join(lost)
            else:
                skills.append({"label": x["label"], "value": ", ".join(lost)})
    return {"summary": summary, "skills": skills, "roles": roles}, added, reverted


# ---------- writing the Word CV and the PDF ----------

def _set_text(paragraph, text: str) -> None:
    runs = paragraph.runs
    runs[0].text = text
    for r in runs[1:]:
        r._r.getparent().remove(r._r)


def _set_skill(paragraph, label: str, value: str) -> None:
    runs = paragraph.runs
    bold = next((r for r in runs if r.bold), runs[0])
    plain = next((r for r in runs if not r.bold), None)
    bold.text = f"{label}: "
    if plain is None:
        plain = paragraph.add_run()
    plain.text = value
    for r in list(paragraph.runs):
        if r._r is not bold._r and r._r is not plain._r:
            r._r.getparent().remove(r._r)


def write_docx(source: Path, content: dict, target: Path) -> None:
    import docx
    document = docx.Document(str(source))
    parsed = parse(document)
    if parsed["summary"] is not None:
        _set_text(parsed["summary"], content["summary"])
    lines = parsed["skills"]
    # Match the number of skill lines: clone the last one or drop extras.
    while len(lines) < len(content["skills"]):
        clone = copy.deepcopy(lines[-1]._p)
        lines[-1]._p.addnext(clone)
        lines = parse(document)["skills"]
    for p in lines[len(content["skills"]):]:
        p._p.getparent().remove(p._p)
    for p, x in zip(parse(document)["skills"], content["skills"]):
        _set_skill(p, x["label"], x["value"])
    for role, new in zip(parse(document)["roles"], content["roles"]):
        for p, text in zip(role["bullets"], new["bullets"]):
            _set_text(p, text)
    target.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(target))


def to_pdf(docx_path: Path, pdf_path: Path) -> None:
    """Convert with a private, invisible Word instance (does not touch documents the user has open)."""
    import pythoncom
    import win32com.client
    pythoncom.CoInitialize()
    word = win32com.client.DispatchEx("Word.Application")
    word.Visible = False
    word.DisplayAlerts = 0
    try:
        doc = word.Documents.Open(str(docx_path), ReadOnly=True, AddToRecentFiles=False)
        doc.SaveAs2(str(pdf_path), FileFormat=17)
        doc.Close(False)
    finally:
        word.Quit()
        pythoncom.CoUninitialize()


def ascii_name(name: str) -> str:
    return "".join(c for c in unicodedata.normalize("NFKD", name) if not unicodedata.combining(c))


def build(job: dict, cv: dict, cfg: dict, language: str, as_email: bool, folder: str | None = None) -> dict:
    """Tailored PDF + cover letter for one offer. Returns paths, added skills, reverted notes and the letter."""
    import docx
    source = Path(cv["docx"])
    facts = base_facts(parse(docx.Document(str(source))))
    proposal = ask_claude(job, facts, cfg, language, as_email)
    content, added, reverted = guard(facts, proposal, cfg)
    target_dir = OUT / (folder or str(job["id"]))
    stem = ascii_name(Path(cv["path"]).stem)
    docx_out, pdf_out = target_dir / f"{stem}.docx", target_dir / f"{stem}.pdf"
    write_docx(source, content, docx_out)
    to_pdf(docx_out, pdf_out)
    return {"pdf": pdf_out, "docx": docx_out, "added": added, "reverted": reverted,
            "missing": proposal["missing"], "cover_letter": proposal["cover_letter"].strip(),
            "summary": content["summary"], "base_summary": facts["summary"]}
